"""可重現的合成日報與修正版教學資料。"""

from datetime import date, timedelta
import random

from openpyxl import Workbook

from .core import FIELDS, QualityError, atomic_workbook, settings


def write_report(path, day, inputs, defects, revision=1, note=""):
    if path.exists():
        raise QualityError(f"不覆蓋既有原始日報：{path.name}")
    wb = Workbook()
    wb.remove(wb.active)
    rows = {
        "日報資訊": [[day, revision, "FULL_DAY", True, note, "SYNTHETIC"]],
        "投入紀錄": inputs,
        "缺陷紀錄": defects,
        "代碼表": [[c, *v] for c, v in settings()["defect_codes"].items()],
    }
    for name, fields in FIELDS.items():
        ws = wb.create_sheet(name)
        ws.append(fields)
        for row in rows[name]:
            ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
    ws = wb.create_sheet("說明")
    ws.append(["合成資料（Synthetic Data）", "非真實工廠經驗；不證明根因"])
    ws.append(["填寫規則", "一天一份完整快照；TRUE 僅代表當日完整；投入不可空白，零缺陷保留標題"])
    ws.append(["正規化", "所有文字 trim；僅 defect_code 統一大寫；日期真日期或 YYYY-MM-DD；投入與版本正整數"])
    ws.append(["修正版", "版本須提高，填原因，仍含全日資料；不得只提供異動"])
    for name, fields in FIELDS.items():
        for field in fields:
            definitions = {
                "report_date": "當日報表日期", "prod_date": "生產日期，等於日報日期",
                "revision_no": "正整數版本", "coverage_mode": "固定 FULL_DAY",
                "is_complete": "布林 TRUE，僅證實當日完整", "revision_note": "修正版必填原因",
                "data_type": "固定 SYNTHETIC", "line": "產線識別", "shift": "早或晚",
                "product": "產品識別", "input_boards": "投入板數，正整數",
                "defect_id": "現行資料全域唯一缺陷編號", "board_sn": "全域板號，僅屬於一組日期產線班別產品",
                "defect_code": "受控主檔缺陷代碼", "position": "合成板上位置",
                "inspection_station": "合成檢驗站別", "defect_name": "受控缺陷名稱",
                "defect_category": "受控缺陷類別",
            }
            ws.append([f"{name}.{field}", definitions[field]])
    atomic_workbook(wb, path, [*FIELDS, "說明"])


def case_rows():
    d1, d2 = date(2026, 10, 5), date(2026, 10, 6)
    inputs = [
        [d1, "L1", "早", "PCB-A", 100], [d1, "L1", "晚", "PCB-A", 150],
        [d1, "L2", "早", "PCB-B", 40], [d2, "L2", "晚", "PCB-B", 300],
    ]
    mapping = [(0, "SN-A001", "SHORT"), (0, "SN-A001", "INSUF"), (0, "SN-A002", "SHORT"),
               (0, "SN-A003", "SHORT"), (1, "SN-A101", "SHIFT"), (2, "SN-B001", "SHORT"),
               (2, "SN-B001", "MISSING"), (2, "SN-B002", "SHIFT"),
               (3, "SN-B201", "INSUF"), (3, "SN-B202", "INSUF")]
    defects = [[f"D{i:03d}", *inputs[p][:4], sn, code, f"U{i}", "外觀檢驗站"]
               for i, (p, sn, code) in enumerate(mapping, 1)]
    return inputs, defects


def generate(mode, base):
    raw, revisions = base / "raw", base / "revisions"
    raw.mkdir(parents=True, exist_ok=True)
    revisions.mkdir(parents=True, exist_ok=True)
    if list(raw.glob("*.xlsx")) or list(revisions.glob("*.xlsx")):
        raise QualityError("已有合成日報，生成器不覆蓋；請保留既有資料或使用新的設定路徑")
    if mode == "case":
        inputs, defects = case_rows()
        for day in sorted({r[0] for r in inputs}):
            write_report(raw / f"品質日報_{day}_v1.xlsx", day,
                         [r for r in inputs if r[0] == day], [r for r in defects if r[1] == day])
        day = date(2026, 10, 5)
        revised_inputs = [r.copy() for r in inputs if r[0] == day]
        revised_inputs[0][-1] = 120
        revised_defects = [r.copy() for r in defects if r[1] == day and r[0] != "D004"]
        revised_defects.append(["D011", day, "L1", "早", "PCB-A", "SN-A004", "MISSING", "U11", "外觀檢驗站"])
        write_report(revisions / f"品質日報_{day}_v2.xlsx", day, revised_inputs, revised_defects,
                     2, "合成教學：投入100改120、刪D004、增D011")
    else:
        rng = random.Random(settings()["seed"])
        codes = list(settings()["defect_codes"])
        for n in range(28):
            day = date(2026, 9, 14) + timedelta(days=n)
            inputs, defects = [], []
            for line in ("L1", "L2", "L3"):
                for shift in ("早", "晚"):
                    for product in ("PCB-A", "PCB-B", "PCB-C"):
                        count = 10 if (line, shift, product) == ("L3", "早", "PCB-C") else rng.randint(110, 180)
                        inputs.append([day, line, shift, product, count])
                        if (line, shift, product) == ("L2", "晚", "PCB-B"):
                            bad = 0
                        elif (line, shift, product) == ("L1", "早", "PCB-A"):
                            bad = 12
                        elif count == 10:
                            bad = 3
                        else:
                            bad = rng.randint(0, 2)
                        for board in range(bad):
                            sn = f"SN-{day}-{line}-{shift}-{product}-{board:03d}"
                            for extra in range(2 if board == 0 else 1):
                                defects.append([f"D-{day}-{len(defects):04d}", day, line, shift, product,
                                                sn, rng.choice(codes), f"U{board+extra+1}", "外觀檢驗站"])
            write_report(raw / f"品質日報_{day}_v1.xlsx", day, inputs, defects)
    return len(list(raw.glob("*.xlsx")))
