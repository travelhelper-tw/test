"""從 SQLite 全量重產 Power BI 用 Excel，不使用文字中間檔。"""

from datetime import date, datetime, timedelta, timezone

import pandas as pd
from openpyxl import Workbook

from .core import ROOT, QualityError, analysis_lock, atomic_workbook, iso_week, settings


def query(db, name, params=None):
    return pd.read_sql_query((ROOT / "sql" / name).read_text(encoding="utf-8"), db, params=params)


def frames(db):
    owns_transaction = not db.in_transaction
    if owns_transaction:
        db.execute("BEGIN")
    try:
        return snapshot_frames(db)
    finally:
        if owns_transaction:
            db.rollback()


def snapshot_frames(db):
    cfg = settings()
    production = query(db, "daily_summary.sql")
    if production.empty:
        raise QualityError("資料庫沒有現行投入，不能輸出分析")
    defects = pd.read_sql_query("SELECT * FROM defect_record ORDER BY prod_date,defect_id", db)
    weekly = query(db, "weekly_summary.sql", {k: cfg[k] for k in ("min_input", "min_bad", "rate_multiplier")})
    counts = query(db, "defect_analysis.sql")
    active_days = {r[0] for r in db.execute("SELECT report_date FROM active_daily_report")}
    for field in ["main_defect", "reason", "rule_version", "provisional_status"]:
        weekly[field] = ""
    weekly["main_defect_count"] = 0
    weekly["main_defect_share"] = 0.0
    weekly["main_defect_tied"] = False
    for index, row in weekly.iterrows():
        group = counts
        for key in ["week_key", "line", "shift", "product"]:
            group = group[group[key] == row[key]]
        if not group.empty:
            top = group[group["defect_count"] == group["defect_count"].max()].sort_values("defect_code")
            weekly.loc[index, "main_defect"] = "、".join(top["defect_name"])
            weekly.loc[index, "main_defect_count"] = int(top["defect_count"].iloc[0])
            weekly.loc[index, "main_defect_share"] = int(top["defect_count"].iloc[0]) / row["defect_count"]
            weekly.loc[index, "main_defect_tied"] = len(top) > 1
        else:
            weekly.loc[index, "main_defect"] = "無缺陷"
        weekly.loc[index, "reason"] = {
            "優先調查": "達同週整體倍數及樣本門檻；核對原始日報、位置與檢驗紀錄，不推論根因",
            "觀察（樣本不足）": "達倍數門檻但投入或不良板樣本不足；補充資料再核查",
            "未觸發規則": "未達相對門檻或全週零不良；不代表合格",
        }[row["priority"]]
        weekly.loc[index, "rule_version"] = cfg["rule_version"]
        year, week = row["week_key"].split("-W")
        monday = date.fromisocalendar(int(year), int(week), 1)
        complete = all((monday + timedelta(days=i)).isoformat() in active_days for i in range(7))
        weekly.loc[index, "provisional_status"] = "七日日報完整" if complete else "截至資料截止日的暫定評估"
    dates = pd.date_range(production["prod_date"].min(), production["prod_date"].max())
    dim_date = pd.DataFrame({"prod_date": [d.date() for d in dates], "week_key": [iso_week(d.date()) for d in dates]})
    week_rows = []
    for key in sorted(dim_date["week_key"].unique()):
        year, week = key.split("-W")
        start = date.fromisocalendar(int(year), int(week), 1)
        week_rows.append([key, start, start + timedelta(days=6)])
    for frame in [production, defects]:
        frame["prod_date"] = pd.to_datetime(frame["prod_date"]).dt.date
    info = pd.DataFrame([
        ["資料聲明", "合成資料（Synthetic Data）；非真實工廠經驗或根因證明"],
        ["匯出時間", datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)],
        ["資料截止日", max(dim_date["prod_date"])],
        ["規則版本", cfg["rule_version"]],
        ["限制", "非統計檢定；未觸發≠合格；不同產品基準與全廠同步惡化可能影響判斷；不推論因果"],
        ["主要缺陷", "並列時筆數與占比為每個並列代碼，不是合計；零缺陷占比為0"],
        ["暫定評估", "每週七個完整日報皆存在才標示完整；同週整體基準不受報表點選重算"],
    ], columns=["項目", "內容"])
    result = {
        "fact_production": production, "fact_defect": defects, "weekly_priority": weekly,
        "dim_week": pd.DataFrame(week_rows, columns=["week_key", "week_start", "week_end"]),
        "dim_date": dim_date,
        "dim_defect": pd.read_sql_query("SELECT * FROM defect_code ORDER BY defect_code", db),
        "說明": info,
    }
    for key in ("line", "shift", "product"):
        result[f"dim_{key}"] = pd.DataFrame({key: sorted(production[key].unique())})
    for name, key in [("dim_date", "prod_date"), ("dim_week", "week_key"),
                      ("dim_line", "line"), ("dim_shift", "shift"),
                      ("dim_product", "product"), ("dim_defect", "defect_code")]:
        dimension = result[name]
        if dimension[key].isna().any() or dimension[key].duplicated().any():
            raise QualityError(f"維度鍵空白或重複：{name}")
        for fact in ["fact_production", "fact_defect", "weekly_priority", "dim_date"]:
            if key in result[fact] and not set(result[fact][key]).issubset(set(dimension[key])):
                raise QualityError(f"資料鍵沒有對應維度：{fact}.{key}")
    return result


def export(db, path):
    with analysis_lock(db):
        return export_locked(db, path)


def export_locked(db, path):
    data = frames(db)
    wb = Workbook()
    wb.remove(wb.active)
    for name, frame in data.items():
        ws = wb.create_sheet(name)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            ws.append([None if pd.isna(value) else value for value in row])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss" if isinstance(cell.value, datetime) else "yyyy-mm-dd"
                if cell.column <= len(frame.columns) and frame.columns[cell.column-1] in ("bad_rate", "baseline_rate", "main_defect_share"):
                    cell.number_format = "0.00%"
    atomic_workbook(wb, path, data.keys())
    return data
