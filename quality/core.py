"""檢查日報、整日替換與可追溯封存。"""

import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sqlite3
from datetime import date, datetime, timezone
from uuid import uuid4

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIELDS = {
    "日報資訊": ["report_date", "revision_no", "coverage_mode", "is_complete", "revision_note", "data_type"],
    "投入紀錄": ["prod_date", "line", "shift", "product", "input_boards"],
    "缺陷紀錄": ["defect_id", "prod_date", "line", "shift", "product", "board_sn", "defect_code", "position", "inspection_station"],
    "代碼表": ["defect_code", "defect_name", "defect_category"],
}
KEY = ["prod_date", "line", "shift", "product"]


class QualityError(Exception):
    """阻擋匯入或輸出的錯誤。"""


def settings():
    return json.loads((ROOT / "config/settings.json").read_text(encoding="utf-8"))


def mode_path(mode):
    return ROOT / settings()["modes"][mode]


def iso_week(value):
    y, w, _ = date.fromisoformat(str(value)[:10]).isocalendar()
    return f"{y}-W{w:02d}"


def connect(path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA foreign_keys=ON")
    db.create_function("iso_week", 1, iso_week, deterministic=True)
    db.executescript((ROOT / "sql/schema.sql").read_text(encoding="utf-8"))
    for code, (name, category) in settings()["defect_codes"].items():
        db.execute("INSERT OR IGNORE INTO defect_code VALUES(?,?,?)", (code, name, category))
        existing = db.execute("SELECT defect_name,defect_category FROM defect_code WHERE defect_code=?", (code,)).fetchone()
        if tuple(existing) != (name, category):
            db.close()
            raise QualityError("資料庫代碼主檔與設定不一致，請先審核設定變更")
    db.commit()
    return db


def digest(path):
    with open(path, "rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def atomic_workbook(workbook, path, required):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        for ws in workbook:
            for row in ws:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        workbook.save(temp)
        check = load_workbook(temp, read_only=True)
        try:
            if not set(required).issubset(check.sheetnames):
                raise QualityError("暫存 Excel 缺少必要工作表")
            def equivalent(expected, actual):
                if expected == "":
                    expected = None
                if isinstance(expected, date) and not isinstance(expected, datetime):
                    expected = datetime.combine(expected, datetime.min.time())
                if isinstance(expected, float) and isinstance(actual, (float, int)):
                    return math.isclose(expected, actual, rel_tol=1e-14, abs_tol=1e-14)
                return expected == actual
            for name in required:
                before, after = list(workbook[name].values), list(check[name].values)
                if len(before) != len(after) or any(
                    len(a) != len(b) or not all(equivalent(x, y) for x, y in zip(a, b))
                    for a, b in zip(before, after)
                ):
                    raise QualityError(f"暫存 Excel 回讀內容不一致：{name}")
        finally:
            check.close()
        os.replace(temp, path)
    except Exception as exc:
        raise QualityError(f"Excel 輸出失敗；舊檔保留，請關閉占用檔案後重試：{exc}") from exc
    finally:
        temp.unlink(missing_ok=True)


def validate(path, db, report_path):
    issues, tables = [], {}

    def issue(sheet, row, field, reason, blocking=True):
        issues.append([sheet, row, field, reason, "未刪除紀錄；阻擋匯入" if blocking else "已正規化", blocking])

    try:
        with open(path, "rb") as stream:
            wb = load_workbook(stream, data_only=False)
    except Exception as exc:
        issue("檔案", 0, "", f"無法讀取 Excel：{exc}")
        wb = None
    if wb:
        try:
            for sheet in [*FIELDS, "說明"]:
                if sheet not in wb.sheetnames:
                    issue(sheet, 1, "", "缺少必要工作表")
            for sheet, fields in FIELDS.items():
                tables[sheet] = []
                if sheet not in wb.sheetnames:
                    continue
                rows = list(wb[sheet].iter_rows(values_only=True))
                headers = [str(v).strip() if v is not None else "" for v in (rows[0] if rows else [])]
                for field in fields:
                    if headers.count(field) != 1:
                        issue(sheet, 1, field, "必要欄位缺少或重複")
                if any(headers.count(f) != 1 for f in fields):
                    continue
                for row_no, values in enumerate(rows[1:], 2):
                    if all(v is None for v in values):
                        issue(sheet, row_no, "", "空白資料列不可默默刪除")
                        continue
                    record = {"source_sheet": sheet, "source_row": row_no}
                    for field in fields:
                        v = values[headers.index(field)]
                        if isinstance(v, str):
                            if v.startswith("="):
                                issue(sheet, row_no, field, "不接受公式，請填寫明確值")
                            clean = v.strip()
                            if field == "defect_code":
                                clean = clean.upper()
                            if clean != v:
                                issue(sheet, row_no, field, "去除前後空白；缺陷代碼轉大寫", False)
                            v = clean
                        if v is None or v == "":
                            if field != "revision_note":
                                issue(sheet, row_no, field, "必填值空白")
                        if field in ("report_date", "prod_date"):
                            try:
                                if isinstance(v, datetime):
                                    if v.time() != datetime.min.time():
                                        raise ValueError()
                                    v = v.date().isoformat()
                                elif isinstance(v, date):
                                    v = v.isoformat()
                                elif isinstance(v, str) and date.fromisoformat(v).isoformat() == v:
                                    pass
                                else:
                                    raise ValueError()
                            except (ValueError, TypeError):
                                issue(sheet, row_no, field, "須為真日期或 YYYY-MM-DD；不可含時間")
                        if field in ("revision_no", "input_boards"):
                            if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) or v <= 0 or int(v) != v:
                                issue(sheet, row_no, field, "須為正整數數值")
                            else:
                                v = int(v)
                        record[field] = v
                    tables[sheet].append(record)
            meta_rows = tables.get("日報資訊", [])
            meta = meta_rows[0] if len(meta_rows) == 1 else {}
            if len(meta_rows) != 1:
                issue("日報資訊", 2, "", "日報資訊必須恰好一列")
            for field, expected in [("coverage_mode", "FULL_DAY"), ("is_complete", True), ("data_type", "SYNTHETIC")]:
                if meta.get(field) != expected or (field == "is_complete" and meta.get(field) is not True):
                    issue("日報資訊", 2, field, f"必須為 {expected}")
            if isinstance(meta.get("revision_no"), int) and meta["revision_no"] > 1 and not meta.get("revision_note"):
                issue("日報資訊", 2, "revision_note", "修正版必填原因")
            inputs, defects = tables.get("投入紀錄", []), tables.get("缺陷紀錄", [])
            if not inputs:
                issue("投入紀錄", 2, "", "投入紀錄不可空白")
            controlled = settings()["defect_codes"]
            seen_codes = set()
            for r in tables.get("代碼表", []):
                c = r["defect_code"]
                if c in seen_codes or controlled.get(c) != [r["defect_name"], r["defect_category"]]:
                    issue("代碼表", r["source_row"], "defect_code", "代碼重複、未知或意義與受控主檔不一致")
                seen_codes.add(c)
            if seen_codes != set(controlled):
                issue("代碼表", 1, "defect_code", "須完整列出受控主檔")
            input_map, ids, boards, bad = {}, set(), {}, {}
            for sheet, records in [("投入紀錄", inputs), ("缺陷紀錄", defects)]:
                for r in records:
                    k = tuple(r[f] for f in KEY)
                    if r["prod_date"] != meta.get("report_date"):
                        issue(sheet, r["source_row"], "prod_date", "紀錄日期不等於日報日期")
                    if r["shift"] not in ("早", "晚"):
                        issue(sheet, r["source_row"], "shift", "班別僅限早、晚")
                    if sheet == "投入紀錄":
                        if k in input_map:
                            issue(sheet, r["source_row"], "", "投入組合重複")
                        input_map[k] = r["input_boards"]
                    else:
                        if k not in input_map:
                            issue(sheet, r["source_row"], "", "缺陷無對應投入")
                        if r["defect_id"] in ids:
                            issue(sheet, r["source_row"], "defect_id", "缺陷編號重複")
                        ids.add(r["defect_id"])
                        if r["defect_code"] not in controlled:
                            issue(sheet, r["source_row"], "defect_code", "未知缺陷代碼")
                        sn = r["board_sn"]
                        if sn in boards and boards[sn] != k:
                            issue(sheet, r["source_row"], "board_sn", "板號跨組合衝突")
                        boards[sn] = k
                        bad.setdefault(k, set()).add(sn)
            for k, sns in bad.items():
                n = input_map.get(k)
                if isinstance(n, int) and len(sns) > n:
                    issue("投入紀錄", 0, "input_boards", "不良板數超過投入")
            for old in db.execute("SELECT * FROM defect_record WHERE prod_date<>?", (meta.get("report_date"),)):
                for r in defects:
                    if old["defect_id"] == r["defect_id"]:
                        issue("缺陷紀錄", r["source_row"], "defect_id", "跨日缺陷編號衝突")
                    if old["board_sn"] == r["board_sn"] and tuple(old[f] for f in KEY) != tuple(r[f] for f in KEY):
                        issue("缺陷紀錄", r["source_row"], "board_sn", "跨日板號衝突")
        finally:
            wb.close()
    report = Workbook()
    ws = report.active
    ws.title = "檢查結果"
    ws.append(["來源工作表", "Excel列號", "錯誤欄位", "原因", "處理摘要", "是否阻擋"])
    for row in issues:
        ws.append(row)
    if not any(r[-1] for r in issues):
        ws.append(["全部工作表", 0, "", "檢查通過", "未刪除紀錄；可進入版本與封存檢查", False])
    report.create_sheet("說明").append(["合成資料（Synthetic Data）", "完整檢查；無疑義紀錄才匯入"])
    atomic_workbook(report, report_path, ["檢查結果", "說明"])
    if any(r[-1] for r in issues):
        raise QualityError("檢查有阻擋性錯誤，請查看檢查報告.xlsx")
    return meta, inputs, defects


def change_summary(db, day, inputs, defects):
    result = {}
    for table, new, fields in [("production_input", inputs, KEY), ("defect_record", defects, ["defect_id"])]:
        old = {tuple(r[f] for f in fields): dict(r) for r in db.execute(f"SELECT * FROM {table} WHERE prod_date=?", (day,))}
        current = {tuple(r[f] for f in fields): r for r in new}
        modified = sum(any(old[k].get(f) != v for f, v in current[k].items() if f not in ("source_row", "source_sheet")) for k in old.keys() & current.keys())
        result[table] = {"新增": len(current.keys() - old.keys()), "修改": modified, "刪除": len(old.keys() - current.keys())}
    old_inputs = db.execute("SELECT COALESCE(SUM(input_boards),0) FROM production_input WHERE prod_date=?", (day,)).fetchone()[0]
    old_bad, old_def = db.execute("SELECT COUNT(DISTINCT board_sn),COUNT(*) FROM defect_record WHERE prod_date=?", (day,)).fetchone()
    result["前後比較（投入／不良板／缺陷）"] = {
        "之前": [old_inputs, old_bad, old_def],
        "之後": [sum(r["input_boards"] for r in inputs), len({r["board_sn"] for r in defects}), len(defects)],
    }
    return result


def import_report(db, path, archive, report_path, confirm=None):
    path, archive = Path(path), Path(archive)
    archive.mkdir(parents=True, exist_ok=True)
    lock = archive / ".import.lock"
    try:
        handle = lock.open("x", encoding="utf-8")
    except FileExistsError as exc:
        raise QualityError("匯入鎖存在；可能另有匯入或中斷殘留，請先依交付指南恢復") from exc
    handle.close()
    token = uuid4().hex
    temp = archive / f".{token}.pending"
    backup = archive / f".{token}.db.pending"
    final = None
    committed = False
    preserve = False
    try:
        if list(archive.glob("*.pending")):
            raise QualityError("存在中斷封存／備份，請先依交付指南恢復，不可直接刪除")
        shutil.copyfile(path, temp)
        sha = digest(temp)
        if sha != digest(path):
            raise QualityError("來源檔案複製期間變更，未匯入")
        meta, inputs, defects = validate(temp, db, report_path)
        day, revision = meta["report_date"], meta["revision_no"]
        active = db.execute("SELECT b.* FROM active_daily_report a JOIN import_batch b USING(batch_id) WHERE a.report_date=?", (day,)).fetchone()
        if active and sha == active["sha256"]:
            return "相同檔案，已跳過"
        if active and revision <= active["revision_no"]:
            raise QualityError("同版不同內容或較低版本不可覆蓋；恢復舊內容須提高版本並填原因")
        summary = change_summary(db, day, inputs, defects)
        print(json.dumps(summary, ensure_ascii=False))
        if active and (confirm is None or not confirm(summary)):
            raise QualityError("使用者未確認，未替換")
        final = archive / f"品質日報_{day}_v{revision}_{sha}.xlsx"
        if final.exists():
            raise QualityError("封存路徑已存在，請先查核殘留檔，不覆蓋")
        snapshot = sqlite3.connect(backup)
        try:
            db.backup(snapshot)
        finally:
            snapshot.close()
        db.execute("BEGIN IMMEDIATE")
        db.execute("DELETE FROM defect_record WHERE prod_date=?", (day,))
        db.execute("DELETE FROM production_input WHERE prod_date=?", (day,))
        for table, rows, fields in [("production_input", inputs, FIELDS["投入紀錄"]), ("defect_record", defects, FIELDS["缺陷紀錄"])]:
            columns = fields + ["batch_id", "source_sheet", "source_row"]
            db.executemany(f"INSERT INTO {table} ({','.join(columns)}) VALUES ({','.join('?' for _ in columns)})",
                           [tuple(r[f] for f in fields) + (token, r["source_sheet"], r["source_row"]) for r in rows])
        db.execute("INSERT INTO import_batch VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (
            token, day, revision, sha, path.name, final.name, datetime.now(timezone.utc).isoformat(),
            meta.get("revision_note") or "初版完整日報", len(inputs), sum(r["input_boards"] for r in inputs),
            len({r["board_sn"] for r in defects}), len(defects)))
        db.execute("INSERT INTO active_daily_report VALUES(?,?) ON CONFLICT(report_date) DO UPDATE SET batch_id=excluded.batch_id", (day, token))
        db.commit()
        committed = True
        try:
            os.rename(temp, final)
        except Exception:
            preserve = True
            restore = sqlite3.connect(backup)
            try:
                restore.backup(db)
            finally:
                restore.close()
            committed = False
            preserve = False
            raise
        return "成功匯入；合成資料（Synthetic Data）"
    except Exception as exc:
        if not committed:
            db.rollback()
        if isinstance(exc, QualityError):
            raise
        message = "恢復失敗，保留鎖、封存與備份，禁止繼續使用，請依交付指南恢復" if preserve else "匯入失敗，未留下成功批次或已回復資料庫"
        raise QualityError(f"{message}：{exc}") from exc
    finally:
        if not preserve:
            temp.unlink(missing_ok=True)
            backup.unlink(missing_ok=True)
            lock.unlink(missing_ok=True)
