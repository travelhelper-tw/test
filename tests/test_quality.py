"""使用獨立暫存資料庫驗證合成案例與失敗回復。"""

from datetime import date
import shutil

from openpyxl import load_workbook
import pytest

from quality import core
from quality.analysis import export, frames
from quality.core import QualityError, connect, import_report, iso_week
from quality.generate import case_rows, generate, write_report


@pytest.fixture
def case(tmp_path):
    generate("case", tmp_path)
    db = connect(tmp_path / "quality.sqlite3")
    yield tmp_path, db
    db.close()


def ingest(case, filename="品質日報_2026-10-05_v1.xlsx", revised=False):
    base, db = case
    return import_report(db, base / ("revisions" if revised else "raw") / filename,
                         base / "archive", base / "output/檢查報告.xlsx", lambda _: True)


def ingest_all(case):
    ingest(case)
    ingest(case, "品質日報_2026-10-06_v1.xlsx")


def edit(case, mutate, revision=False):
    base, _ = case
    path = base / ("revisions/品質日報_2026-10-05_v2.xlsx" if revision else "raw/品質日報_2026-10-05_v1.xlsx")
    wb = load_workbook(path)
    mutate(wb)
    wb.save(path)
    wb.close()


@pytest.mark.parametrize("mutate", [
    lambda w: w.remove(w["缺陷紀錄"]),
    lambda w: w.remove(w["說明"]),
    lambda w: setattr(w["投入紀錄"]["E1"], "value", "錯誤欄位"),
    lambda w: setattr(w["投入紀錄"]["B2"], "value", " "),
    lambda w: setattr(w["投入紀錄"]["A2"], "value", "2026-02-30"),
    lambda w: (setattr(w["投入紀錄"]["A2"], "value", 46300), setattr(w["投入紀錄"]["A2"], "number_format", "General")),
    lambda w: setattr(w["投入紀錄"]["A2"], "value", date(2026, 10, 6)),
    lambda w: setattr(w["缺陷紀錄"]["G2"], "value", "UNKNOWN"),
    lambda w: setattr(w["代碼表"]["B2"], "value", "擅改意義"),
    lambda w: setattr(w["投入紀錄"]["E2"], "value", 0),
    lambda w: setattr(w["投入紀錄"]["E2"], "value", 1.5),
    lambda w: setattr(w["投入紀錄"]["E2"], "value", "100"),
    lambda w: setattr(w["投入紀錄"]["C2"], "value", "中"),
    lambda w: w["投入紀錄"].append([c.value for c in w["投入紀錄"][2]]),
    lambda w: setattr(w["缺陷紀錄"]["A3"], "value", "D001"),
    lambda w: setattr(w["缺陷紀錄"]["F6"], "value", "SN-A001"),
    lambda w: setattr(w["缺陷紀錄"]["C2"], "value", "L9"),
    lambda w: setattr(w["投入紀錄"]["E2"], "value", 2),
    lambda w: setattr(w["日報資訊"]["D2"], "value", False),
    lambda w: setattr(w["日報資訊"]["F2"], "value", "REAL"),
    lambda w: w["投入紀錄"].append([None, None, None, None, ""]),
])
def test_invalid_workbook_blocked(case, mutate):
    edit(case, mutate)
    with pytest.raises(QualityError, match="阻擋"):
        ingest(case)
    base, db = case
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == 0
    assert not list((base / "archive").iterdir())
    wb = load_workbook(base / "output/檢查報告.xlsx")
    assert any(row[-1] is True for row in wb["檢查結果"].iter_rows(min_row=2, values_only=True))
    wb.close()


def test_trim_and_upper_preserve_lineage(case):
    edit(case, lambda w: setattr(w["缺陷紀錄"]["G2"], "value", " short "))
    ingest(case)
    base, db = case
    r = db.execute("SELECT * FROM defect_record WHERE defect_id='D001'").fetchone()
    assert (r["defect_code"], r["source_sheet"], r["source_row"]) == ("SHORT", "缺陷紀錄", 2)
    wb = load_workbook(base / "output/檢查報告.xlsx")
    assert any(row[-1] is False for row in wb["檢查結果"].iter_rows(min_row=2, values_only=True))
    wb.close()


def test_manual_case_and_export(case):
    ingest_all(case)
    base, db = case
    data = export(db, base / "output/品質分析.xlsx")
    p, w = data["fact_production"], data["weekly_priority"]
    assert list(p[["input_boards", "bad_boards", "defect_count"]].sum()) == [590, 8, 10]
    assert set(w.baseline_rate) == {8/590}
    assert w.priority.tolist() == ["優先調查", "未觸發規則", "觀察（樣本不足）", "未觸發規則"]
    assert w.bad_rate.tolist() == [3/100, 1/150, 2/40, 2/300]
    first = w.iloc[0]
    assert (first.main_defect, first.main_defect_count, first.main_defect_share) == ("短路", 3, .75)
    assert set(w.provisional_status) == {"截至資料截止日的暫定評估"}
    assert data["fact_defect"].groupby("defect_code").size().to_dict() == {"SHORT": 4, "INSUF": 3, "SHIFT": 2, "MISSING": 1}
    l1 = p[p.line == "L1"]
    assert list(l1[["input_boards", "bad_boards", "defect_count"]].sum()) == [250, 4, 5]
    assert list(l1[l1["shift"] == "早"][["input_boards", "bad_boards", "defect_count"]].sum()) == [100, 3, 4]
    wb = load_workbook(base / "output/品質分析.xlsx")
    assert set(wb.sheetnames) == set(data)
    assert wb["fact_production"]["A2"].is_date
    assert wb["dim_week"]["B2"].is_date
    wb.close()
    assert db.execute("PRAGMA foreign_key_check").fetchall() == []


def test_identical_import_skipped(case):
    ingest(case)
    base, db = case
    before = db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0]
    assert "跳過" in ingest(case)
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == before
    assert len(list((base / "archive").glob("*.xlsx"))) == 1


def test_revision_and_other_day_unchanged(case):
    ingest_all(case)
    base, db = case
    old_other = [tuple(r) for r in db.execute("SELECT * FROM production_input WHERE prod_date='2026-10-06'")]
    old_defects = [tuple(r) for r in db.execute("SELECT * FROM defect_record WHERE prod_date='2026-10-06'")]
    ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    assert [tuple(r) for r in db.execute("SELECT * FROM production_input WHERE prod_date='2026-10-06'")] == old_other
    assert [tuple(r) for r in db.execute("SELECT * FROM defect_record WHERE prod_date='2026-10-06'")] == old_defects
    w = frames(db)["weekly_priority"]
    assert list(w[["input_boards", "bad_boards", "defect_count"]].sum()) == [610, 8, 10]
    assert w.iloc[0].priority == "未觸發規則"
    assert w.iloc[0].main_defect_share == .5
    assert db.execute("SELECT COUNT(*) FROM defect_record WHERE defect_id='D004'").fetchone()[0] == 0
    assert db.execute("SELECT COUNT(*) FROM defect_record WHERE defect_id='D011'").fetchone()[0] == 1
    assert len(list((base / "archive").glob("*.xlsx"))) == 3
    assert db.execute("PRAGMA foreign_key_check").fetchall() == []


def test_revision_delete_whole_combination(case):
    ingest_all(case)
    def mutate(w):
        w["投入紀錄"].delete_rows(4)
        w["缺陷紀錄"].delete_rows(6, 3)
    edit(case, mutate, True)
    ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    _, db = case
    assert db.execute("SELECT COUNT(*) FROM production_input WHERE prod_date='2026-10-05'").fetchone()[0] == 2
    assert db.execute("SELECT COUNT(*) FROM defect_record WHERE line='L2' AND prod_date='2026-10-05'").fetchone()[0] == 0


def test_same_revision_changed_and_old_rejected(case):
    ingest(case)
    edit(case, lambda w: setattr(w["投入紀錄"]["E2"], "value", 101))
    with pytest.raises(QualityError, match="同版"):
        ingest(case)
    ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    with pytest.raises(QualityError, match="較低"):
        ingest(case)


def test_invalid_revision_keeps_old(case):
    ingest_all(case)
    _, db = case
    before = db_dump(db)
    edit(case, lambda w: setattr(w["日報資訊"]["E2"], "value", ""), True)
    with pytest.raises(QualityError):
        ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    assert db_dump(db) == before


def db_dump(db):
    return "\n".join(db.iterdump())


def test_cross_day_conflicts(case):
    ingest(case)
    base, _ = case
    path = base / "raw/品質日報_2026-10-06_v1.xlsx"
    wb = load_workbook(path)
    wb["缺陷紀錄"]["A2"] = "D001"
    wb["缺陷紀錄"]["F3"] = "SN-A001"
    wb.save(path)
    wb.close()
    with pytest.raises(QualityError):
        ingest(case, path.name)


def test_transaction_failure_rolls_back_and_cleans_archive(case):
    ingest_all(case)
    base, db = case
    db.execute("CREATE TRIGGER 故障 BEFORE INSERT ON import_batch WHEN NEW.revision_no=2 BEGIN SELECT RAISE(ABORT,'合成故障'); END")
    db.commit()
    before = db_dump(db)
    with pytest.raises(QualityError, match="匯入失敗"):
        ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    assert db_dump(db) == before
    assert len(list((base / "archive").iterdir())) == 2


def test_archive_rename_failure_compensates(case, monkeypatch):
    ingest_all(case)
    base, db = case
    before = db_dump(db)
    def fail(*args):
        raise PermissionError("合成封存失敗")
    monkeypatch.setattr(core.os, "rename", fail)
    with pytest.raises(QualityError, match="已回復"):
        ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    assert db_dump(db) == before
    assert len(list((base / "archive").iterdir())) == 2


def test_archive_copy_failure_no_batch(case, monkeypatch):
    base, db = case
    def fail(*args):
        raise OSError("合成複製失敗")
    monkeypatch.setattr(shutil, "copyfile", fail)
    with pytest.raises(QualityError):
        ingest(case)
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == 0
    assert not list((base / "archive").iterdir())


def test_revision_requires_confirmation(case):
    ingest(case)
    base, db = case
    with pytest.raises(QualityError, match="未確認"):
        import_report(db, base / "revisions/品質日報_2026-10-05_v2.xlsx", base / "archive", base / "check.xlsx", lambda _: False)
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == 1


def test_zero_defect_week_and_zero_rows(tmp_path):
    day = date(2026, 10, 5)
    path = tmp_path / "zero.xlsx"
    write_report(path, day, [[day, "L1", "早", "PCB-A", 100]], [])
    db = connect(tmp_path / "db.sqlite3")
    try:
        import_report(db, path, tmp_path / "archive", tmp_path / "check.xlsx")
        data = export(db, tmp_path / "品質分析.xlsx")
        w = data["weekly_priority"].iloc[0]
        assert w.priority == "未觸發規則"
        assert w.main_defect == "無缺陷"
        assert w.main_defect_share == 0
        assert data["weekly_priority"].rate_multiple.isna().all()
        assert data["fact_defect"].empty
    finally:
        db.close()


def test_tied_main_defects(case):
    def mutate(w):
        w["缺陷紀錄"]["G4"] = "INSUF"
    edit(case, mutate)
    ingest(case)
    _, db = case
    w = frames(db)["weekly_priority"].iloc[0]
    assert w.main_defect_tied
    assert w.main_defect == "錫少、短路"
    assert w.main_defect_count == 2
    assert w.main_defect_share == .5


@pytest.mark.parametrize("day", ["2026-12-31", "2027-01-01"])
def test_cross_year_iso(day):
    assert iso_week(day) == "2026-W53"


def test_cross_year_actual_sql_and_export(tmp_path):
    db = connect(tmp_path / "db.sqlite3")
    try:
        for text in ["2026-12-31", "2027-01-01"]:
            day = date.fromisoformat(text)
            path = tmp_path / f"{day}.xlsx"
            write_report(path, day, [[day, "L1", "早", "PCB-A", 100]], [])
            import_report(db, path, tmp_path / "archive", tmp_path / "check.xlsx")
        data = frames(db)
        assert data["weekly_priority"].week_key.tolist() == ["2026-W53"]
        assert data["weekly_priority"].input_boards.tolist() == [200]
        assert data["dim_week"].iloc[0].week_start == date(2026, 12, 28)
    finally:
        db.close()


def test_missing_whole_week_exports_calendar_dimensions(tmp_path):
    db = connect(tmp_path / "db.sqlite3")
    try:
        for text in ["2026-10-05", "2026-11-02"]:
            day = date.fromisoformat(text)
            path = tmp_path / f"{day}.xlsx"
            write_report(path, day, [[day, "L1", "早", "PCB-A", 100]], [])
            import_report(db, path, tmp_path / "archive", tmp_path / "check.xlsx")
        data = export(db, tmp_path / "品質分析.xlsx")
        assert len(data["dim_week"]) == 5
        assert len(data["weekly_priority"]) == 2
        assert set(data["dim_date"].week_key).issubset(set(data["dim_week"].week_key))
    finally:
        db.close()


def test_excel_replace_failure_keeps_old(case, monkeypatch):
    ingest_all(case)
    base, db = case
    path = base / "output/品質分析.xlsx"
    export(db, path)
    old = path.read_bytes()
    def fail(*args):
        raise PermissionError("合成檔案占用")
    monkeypatch.setattr(core.os, "replace", fail)
    with pytest.raises(QualityError, match="舊檔保留"):
        export(db, path)
    assert path.read_bytes() == old
    assert not list(path.parent.glob("*.tmp.xlsx"))


def test_excel_save_failure_keeps_old(case, monkeypatch):
    ingest_all(case)
    base, db = case
    path = base / "output/品質分析.xlsx"
    export(db, path)
    old = path.read_bytes()
    def fail(*args):
        raise OSError("合成寫入失敗")
    monkeypatch.setattr(core.Workbook, "save", fail)
    with pytest.raises(QualityError):
        export(db, path)
    assert path.read_bytes() == old


def test_sunday_alone_not_complete(tmp_path):
    day = date(2026, 10, 11)
    path = tmp_path / "日報.xlsx"
    write_report(path, day, [[day, "L1", "早", "PCB-A", 100]], [])
    db = connect(tmp_path / "db.sqlite3")
    try:
        import_report(db, path, tmp_path / "archive", tmp_path / "check.xlsx")
        assert frames(db)["weekly_priority"].iloc[0].provisional_status == "截至資料截止日的暫定評估"
    finally:
        db.close()


def test_stale_pending_blocks_import(case):
    base, db = case
    (base / "archive").mkdir()
    pending = base / "archive/.殘留.pending"
    pending.write_bytes(b"")
    with pytest.raises(QualityError, match="中斷"):
        ingest(case)
    assert pending.exists()
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == 0


def test_archive_lock_blocks_second_import(case):
    base, db = case
    (base / "archive").mkdir()
    lock = base / "archive/.import.lock"
    lock.write_text("合成匯入者", encoding="utf-8")
    with pytest.raises(QualityError, match="匯入鎖"):
        ingest(case)
    assert lock.exists()
    assert db.execute("SELECT COUNT(*) FROM import_batch").fetchone()[0] == 0


def test_formula_rejected(case):
    edit(case, lambda w: setattr(w["投入紀錄"]["B2"], "value", "=1+1"))
    with pytest.raises(QualityError):
        ingest(case)


def test_export_text_not_formula(case):
    ingest(case)
    base, db = case
    db.execute("UPDATE defect_record SET position='=1+1'")
    db.commit()
    export(db, base / "output/品質分析.xlsx")
    wb = load_workbook(base / "output/品質分析.xlsx")
    assert wb["fact_defect"]["H2"].data_type == "s"
    wb.close()


def test_recovery_failure_retains_evidence(case, monkeypatch):
    ingest(case)
    base, db = case
    def fail(*args):
        raise PermissionError("合成封存失敗")
    original = core.sqlite3.connect
    calls = []
    def connection(path, *args, **kwargs):
        calls.append(path)
        if len(calls) == 2:
            raise OSError("合成恢復失敗")
        return original(path, *args, **kwargs)
    monkeypatch.setattr(core.os, "rename", fail)
    monkeypatch.setattr(core.sqlite3, "connect", connection)
    with pytest.raises(QualityError, match="恢復失敗"):
        ingest(case, "品質日報_2026-10-05_v2.xlsx", True)
    assert (base / "archive/.import.lock").exists()
    assert len(list((base / "archive").glob("*.pending"))) == 2
    with pytest.raises(QualityError, match="匯入鎖"):
        ingest(case)


def test_package_excludes_environment_and_temporary(tmp_path, monkeypatch):
    from quality import __main__ as cli
    monkeypatch.setattr(cli, "ROOT", tmp_path)
    for relative in ["venv6/套件.txt", "__pycache__/快取.pyc", ".pytest_cache/快取",
                     "data/case/raw/~$測試.xlsx", "data/case/raw/.測試.tmp.xlsx",
                     "data/case/quality.sqlite3", "quality/程式.py", ".gitignore"]:
        path = tmp_path / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"")
    import zipfile
    with zipfile.ZipFile(cli.package()) as archive:
        assert set(archive.namelist()) == {"data/case/quality.sqlite3", "quality/程式.py", ".gitignore"}


def test_demo_reproducible_and_complete(tmp_path):
    first, second = tmp_path / "first", tmp_path / "second"
    generate("demo", first)
    generate("demo", second)
    files = sorted((first / "raw").glob("*.xlsx"))
    assert len(files) == 28
    for a in files:
        wa, wb = load_workbook(a), load_workbook(second / "raw" / a.name)
        assert list(wa["缺陷紀錄"].values) == list(wb["缺陷紀錄"].values)
        assert list(wa["投入紀錄"].values) == list(wb["投入紀錄"].values)
        wa.close()
        wb.close()
    db = connect(tmp_path / "demo.sqlite3")
    try:
        for path in files:
            import_report(db, path, tmp_path / "archive", tmp_path / "check.xlsx")
        data = export(db, tmp_path / "品質分析.xlsx")
        w = data["weekly_priority"]
        assert len(w) == 72
        assert set(w.priority) == {"優先調查", "觀察（樣本不足）", "未觸發規則"}
        assert set(w.provisional_status) == {"七日日報完整"}
        assert (data["fact_production"].defect_count == 0).any()
        assert data["fact_production"].defect_count.sum() > data["fact_production"].bad_boards.sum()
        assert len(data["dim_product"]) == len(data["dim_line"]) == 3
    finally:
        db.close()
